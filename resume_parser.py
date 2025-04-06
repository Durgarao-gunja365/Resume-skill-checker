

import os
os.environ["STREAMLIT_DISABLE_FILE_WATCHING"] = "1"

import streamlit as st
import spacy
from docx import Document
from typing import List, Union
from io import BytesIO
import os
import pandas as pd
import tempfile
from fpdf import FPDF  # For PDF export
from PyPDF2 import PdfReader
from io import BytesIO
# Load spaCy English model

import spacy
from spacy.cli import download

try:
    nlp = spacy.load("en_core_web_sm")
except OSError:
    download("en_core_web_sm")
    nlp = spacy.load("en_core_web_sm")



# ✅ Define known skills
SKILLS_DB = [
     "Python", "Java", "C++", "C", "C#", "JavaScript", "TypeScript", "Go", "Ruby", "PHP", "Swift", "Kotlin", "R", "Rust",

    # Web Development
    "HTML", "CSS", "React", "Angular", "Vue.js", "Next.js", "Svelte", "Tailwind CSS", "Bootstrap", "jQuery",

    # Backend & Frameworks
    "Node.js", "Express.js", "Django", "Flask", "Spring Boot", "Ruby on Rails", ".NET", "Laravel", "FastAPI", "NestJS",

    # Databases
    "SQL", "MySQL", "PostgreSQL", "MongoDB", "Oracle", "SQLite", "MariaDB", "Cassandra", "Redis", "Firebase",

    # Cloud & DevOps
    "AWS", "Azure", "Google Cloud", "Docker", "Kubernetes", "Terraform", "Jenkins", "Ansible", "GitHub Actions", "CI/CD",

    # Tools & Platforms
    "Git", "GitHub", "GitLab", "Bitbucket", "Jira", "Confluence", "Postman", "Swagger", "Figma", "Notion", "Slack",

    # Testing & QA
    "Selenium", "JUnit", "Mockito", "Cypress", "Playwright", "Jest", "Mocha", "Chai", "PyTest", "Postman", "TestNG",

    # Mobile Development
    "Flutter", "React Native", "Android", "iOS", "SwiftUI", "Xcode", "Jetpack Compose",

    # Data Science & AI
    "Pandas", "NumPy", "Matplotlib", "Seaborn", "Scikit-learn", "TensorFlow", "PyTorch", "Keras", "OpenCV", "HuggingFace",

    # Databases & ORM
    "Hibernate", "JPA", "Entity Framework", "Sequelize", "Mongoose", "ORM", "NoSQL",

    # Soft Skills & Concepts
    "Agile", "Scrum", "Kanban", "SDLC", "OOP", "REST API", "GraphQL", "MVC", "Microservices",

    # Cybersecurity & Networking
    "OAuth", "JWT", "HTTPS", "SSL", "Firewalls", "VPN", "Penetration Testing", "Nmap", "Wireshark", "Burp Suite",

    # Analytics & BI
    "Excel", "Power BI", "Tableau", "Looker", "Google Analytics", "Apache Superset",

    # Miscellaneous
    "Linux", "Unix", "Bash", "Zsh", "Shell Scripting", "Makefile", "Nginx", "Apache", "Load Balancer", "Redis", "Memcached"
]
CLEAN_SKILLS_DB = [skill.lower() for skill in SKILLS_DB]

# ✅ Skill aliases
SKILL_ALIASES = {
    "sql": ["mysql", "postgresql", "sqlite"],
    "c++": ["c plus plus"],
    "javascript": ["js"],
    "python": ["py"],
}

def normalize_skill(skill: str) -> str:
    skill = skill.lower()
    for canonical, variants in SKILL_ALIASES.items():
        if skill == canonical or skill in variants:
            return canonical
    return skill

def get_normalized_skills(skills: List[str]) -> List[str]:
    return [normalize_skill(skill) for skill in skills]

# ✅ File extractors

# File extractors (same as your version)
def extract_text_from_docx(file: Union[str, BytesIO]) -> str:
    doc = Document(file)
    return "\n".join([p.text for p in doc.paragraphs])

def extract_text_from_txt(file: Union[str, BytesIO]) -> str:
    if isinstance(file, BytesIO):
        return file.read().decode("utf-8")
    with open(file, "r", encoding="utf-8") as f:
        return f.read()

def extract_text_from_pdf(file: Union[str, BytesIO]) -> str:
    reader = PdfReader(file if isinstance(file, BytesIO) else open(file, "rb"))
    return "\n".join([page.extract_text() for page in reader.pages if page.extract_text()])

# Extract skills
def extract_skills_from_text(text: str) -> List[str]:
    doc = nlp(text.lower())
    found_skills = set()
    words = [token.text for token in doc if not token.is_stop and not token.is_punct]
    for n in range(1, 4):
        ngrams = zip(*[words[i:] for i in range(n)])
        for gram in ngrams:
            phrase = " ".join(gram)
            if phrase in CLEAN_SKILLS_DB:
                found_skills.add(phrase)
    return [skill.title() for skill in found_skills]

# Estimate experience from keywords
def estimate_experience(text: str) -> int:
    import re
    matches = re.findall(r'(\d+)\+?\s+(years|yrs)', text.lower())
    years = sum(int(match[0]) for match in matches)
    return min(years, 30)  # Cap experience

# Generate PDF report

def generate_pdf_report(matched_skills, missing_skills, score, years):
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    pdf.cell(0, 10, "Resume Skill Match Report", ln=1)
    pdf.cell(0, 10, f"Match Score: {score}%", ln=1)
    pdf.cell(0, 10, f"Matched Skills: {', '.join(matched_skills)}", ln=1)
    pdf.cell(0, 10, f"Missing Skills: {', '.join(missing_skills)}", ln=1)
    pdf.cell(0, 10, f"Estimated Years of Experience: {years}", ln=1)

    # Save PDF output to BytesIO buffer
    pdf_output = pdf.output(dest='S').encode('latin1')  # PDF as bytes
    buffer = BytesIO(pdf_output)
    buffer.seek(0)
    return buffer

# Resume processor with enhancements

def generate_ai_suggestions(resume_text, job_description):
    prompt = f"""
    Given this resume:\n{resume_text}\n
    And this job description:\n{job_description}\n
    Suggest:
    1. 3 missing technical skills
    2. Improvements in resume phrasing
    3. Recommended certifications
    4. A tailored summary statement
    """

    response = openai.ChatCompletion.create(
        model="gpt-4",
        messages=[{"role": "user", "content": prompt}],
        temperature=0.7
    )

    return response.choices[0].message["content"]
def process_resume(uploaded_file, filename, job_skills):
    with st.spinner("Analyzing resume..."):
        with tempfile.NamedTemporaryFile(delete=False, suffix=os.path.splitext(filename)[1]) as tmp_file:
            tmp_file.write(uploaded_file.read())
            file_path = tmp_file.name

        if file_path.endswith(".docx"):
            text = extract_text_from_docx(file_path)
        elif file_path.endswith(".txt"):
            text = extract_text_from_txt(file_path)
        elif file_path.endswith(".pdf"):
            text = extract_text_from_pdf(file_path)
        else:
            st.error("Unsupported file format")
            return

        if not text.strip():
            st.warning("The resume appears to be empty.")
            return

        resume_skills = extract_skills_from_text(text)
        resume_skills_normalized = get_normalized_skills(resume_skills)
        job_skills_normalized = get_normalized_skills(job_skills)

        matched_skills = [job_skill for job_skill in job_skills if normalize_skill(job_skill) in resume_skills_normalized]
        missing_skills = [job_skill for job_skill in job_skills if normalize_skill(job_skill) not in resume_skills_normalized]
        score = int((len(matched_skills) / len(job_skills)) * 100) if job_skills else 0
        years_of_experience = estimate_experience(text)

        st.subheader("📄 Extracted Skills")
        st.write(", ".join(resume_skills) if resume_skills else "No skills found.")

        st.success(f"✅ Matched Skills ({len(matched_skills)}):")
        st.write(", ".join(matched_skills))

        st.error(f"❌ Missing Skills ({len(missing_skills)}):")
        st.write(", ".join(missing_skills))

        st.info("📊 Match Score")
        st.metric(label="Match Percentage", value=f"{score}%")

        st.info("🧓 Estimated Experience")
        st.metric(label="Years", value=f"{years_of_experience} years")

        st.download_button(
            label="📥 Download PDF Report",
            data=generate_pdf_report(matched_skills, missing_skills, score, years_of_experience),
            file_name="resume_skill_match_report.pdf",
            mime="application/pdf"
        )

# Streamlit UI
def main():
    st.set_page_config(page_title="Resume Matcher", page_icon="📄")
    st.title("📄 Resume Skill Matcher with Job Description Analyzer")

    uploaded_file = st.file_uploader("📁 Upload your resume (.pdf, .docx, .txt)", type=["pdf", "docx", "txt"])
    job_skill_input = st.text_input("Enter job skills (comma-separated):", "Python, SQL, Django")

    csv_file = st.file_uploader("📄 Or upload a CSV file with skills", type=["csv"])

    job_skills = []
    if csv_file is not None:
        try:
            df = pd.read_csv(csv_file)
            if "skills" in df.columns:
                job_skills = df["skills"].dropna().tolist()
            else:
                st.error("CSV must have a column named 'skills'")
                return
        except Exception as e:
            st.error(f"Error reading CSV: {e}")
            return
    else:
        # Use manually entered skills if CSV is not uploaded
        job_skills = [skill.strip() for skill in job_skill_input.split(",") if skill.strip()]

    if uploaded_file and job_skills:
        process_resume(uploaded_file, uploaded_file.name, job_skills)




if __name__ == "__main__":
    main()
